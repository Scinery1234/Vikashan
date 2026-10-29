"""Word documents for New Page members, built with the Resume-Generator builder.

POST JSON with a member session token:
  {action: 'resume', data}  -> {docx (base64), html}
  {action: 'cover',  data}  -> {docx (base64), html}
  {action: 'extract', filename, file (base64)} -> {text}
  {action: 'ping'} -> {ok: true}   (used by the admin system check)
Stateless: nothing is stored here.
"""
import base64
import hashlib
import hmac
import io
import json
import os
import sys
import time
from http.server import BaseHTTPRequestHandler

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "_py"))

import doc_builder as rb  # noqa: E402  (copied unchanged from Resume-Generator)
import cover_builder as cb  # noqa: E402

MAX_BODY = 4_000_000
STYLE = "classic"  # the one style offered: single column, easy for recruiters' systems to read


def _b64url(s: str) -> bytes:
    return base64.urlsafe_b64decode(s + "=" * (-len(s) % 4))


def verify_session(token: str) -> bool:
    """Same check as lib/newpage-auth.js: HS256, secret derived from JWT_SECRET."""
    try:
        head, body, sig = token.split(".")
        secret = hmac.new((os.environ.get("JWT_SECRET") or "").encode(), b"newpage-member-v1", hashlib.sha256).hexdigest()
        want = hmac.new(secret.encode(), f"{head}.{body}".encode(), hashlib.sha256).digest()
        if not hmac.compare_digest(want, _b64url(sig)):
            return False
        p = json.loads(_b64url(body))
        return p.get("typ") == "session" and float(p.get("exp", 0)) > time.time()
    except Exception:
        return False


def extract_text(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(content))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                cells = []
                for cell in row.cells:
                    txt = " ".join(p.text for p in cell.paragraphs if p.text.strip())
                    if txt and txt not in cells:
                        cells.append(txt)
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(t for t in (pg.extract_text() for pg in reader.pages) if t)
    if name.endswith((".txt", ".md", ".rtf")):
        return content.decode("utf-8", errors="ignore")
    raise ValueError("Upload a Word (.docx), PDF or text file.")


def respond(h, code, obj):
    raw = json.dumps(obj).encode()
    h.send_response(code)
    h.send_header("Content-Type", "application/json")
    h.send_header("Content-Length", str(len(raw)))
    h.end_headers()
    h.wfile.write(raw)


class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        try:
            n = int(self.headers.get("Content-Length") or 0)
            if n > MAX_BODY:
                return respond(self, 413, {"error": "That file is too big. Keep it under 3 MB."})
            b = json.loads(self.rfile.read(n) or b"{}")
        except Exception:
            return respond(self, 400, {"error": "Bad request."})

        if not verify_session(str(b.get("token") or "")):
            return respond(self, 401, {"error": "Please sign in again.", "code": "auth"})

        action = b.get("action")
        template = STYLE
        try:
            if action == "ping":
                return respond(self, 200, {"ok": True})
            if action == "extract":
                text = extract_text(str(b.get("filename") or ""), base64.b64decode(b.get("file") or ""))
                if len(text.strip()) < 40:
                    return respond(self, 422, {"error": "We couldn’t read any text in that file. If it’s a scan or a picture, paste the text instead."})
                return respond(self, 200, {"text": text[:30000]})
            if action == "resume":
                data = b.get("data") or {}
                builder = rb.ResumeBuilder()
                docx = builder.build_word_document_bytes(data, template)
                return respond(self, 200, {"docx": base64.b64encode(docx).decode(), "html": builder.build_html_preview(data, template)})
            if action == "cover":
                data = b.get("data") or {}
                docx = cb.build_cover_docx_bytes(data, template)
                return respond(self, 200, {"docx": base64.b64encode(docx).decode(), "html": cb.build_cover_html(data, template)})
            return respond(self, 400, {"error": "Unknown action."})
        except ValueError as e:
            return respond(self, 422, {"error": str(e)})
        except Exception as e:  # keep the real reason in the Vercel log
            print("docx error:", repr(e), file=sys.stderr)
            return respond(self, 500, {"error": "We couldn’t build that file. Please try again."})

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()
