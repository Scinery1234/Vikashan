"""Cover letter .docx and HTML preview, in the same templates as the resume builder.

Uses the colours, fonts and sizes from doc_builder.TEMPLATES so a resume and its
cover letter look like a matched pair. Single column, one page, Australian
business-letter layout: sender block, date, recipient, subject line, body,
"Yours sincerely".
"""
import html
import io
from typing import Dict

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

import doc_builder as rb


def _clean(c: Dict) -> Dict:
    c = c or {}
    contact = c.get("contact") or {}
    return {
        "name": str(c.get("name") or ""),
        "contact": {k: str(contact.get(k) or "") for k in ("phone", "email", "location", "linkedin")},
        "date": str(c.get("date") or ""),
        "recipient": [str(x) for x in (c.get("recipient") or []) if str(x).strip()],
        "subject": str(c.get("subject") or ""),
        "greeting": str(c.get("greeting") or "Dear Hiring Manager,"),
        "paragraphs": [str(p) for p in (c.get("paragraphs") or []) if str(p).strip()],
        "closing": str(c.get("closing") or "Yours sincerely,"),
        "signoff": str(c.get("signoff") or c.get("name") or ""),
    }


def _contact_line(c: Dict) -> str:
    k = c["contact"]
    return " | ".join(x for x in (k["location"], k["phone"], k["email"], k["linkedin"]) if x)


def _para(doc, text, tmpl, bold=False, size=None, color=None, align=None, after=8, before=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = tmpl["body_font"]
    r.font.size = size or tmpl["body_size"]
    r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    rb._set_para_spacing(p, before=before, after=after)
    return p


def build_cover_docx_bytes(data: Dict, template_id: str = "classic") -> bytes:
    c = _clean(data)
    tmpl = rb._get_template(template_id)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = tmpl["body_font"]
    normal.font.size = Pt(11)
    rb._set_doc_margins(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
    body = Pt(11)

    _para(doc, c["name"], tmpl, bold=True, size=tmpl["name_size"], color=tmpl["heading_color"], after=2)
    line = _contact_line(c)
    if line:
        _para(doc, line, tmpl, size=tmpl["contact_size"], color=tmpl["muted_color"], after=2)
    rb._add_horizontal_rule(doc, tmpl["rule_color"])
    if c["date"]:
        _para(doc, c["date"], tmpl, size=body, after=10, before=8)
    for i, r in enumerate(c["recipient"]):
        _para(doc, r, tmpl, size=body, after=0 if i < len(c["recipient"]) - 1 else 10)
    if c["subject"]:
        _para(doc, c["subject"], tmpl, bold=True, size=body, after=8)
    _para(doc, c["greeting"], tmpl, size=body, after=8)
    for p in c["paragraphs"]:
        _para(doc, p, tmpl, size=body, after=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    _para(doc, c["closing"], tmpl, size=body, after=26, before=4)
    _para(doc, c["signoff"], tmpl, bold=True, size=body, after=0)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_cover_html(data: Dict, template_id: str = "classic") -> str:
    c = _clean(data)
    t = rb._get_template(template_id)
    e = html.escape
    font = "Georgia, serif" if t["body_font"] == "Georgia" else "Calibri, Arial, sans-serif"
    parts = [
        f'<div class="np-page"><h1>{e(c["name"])}</h1>',
        f'<p class="muted">{e(_contact_line(c))}</p><hr>',
    ]
    if c["date"]:
        parts.append(f'<p>{e(c["date"])}</p>')
    if c["recipient"]:
        parts.append("<p>" + "<br>".join(e(x) for x in c["recipient"]) + "</p>")
    if c["subject"]:
        parts.append(f'<p><b>{e(c["subject"])}</b></p>')
    parts.append(f'<p>{e(c["greeting"])}</p>')
    parts += [f"<p>{e(p)}</p>" for p in c["paragraphs"]]
    parts.append(f'<p>{e(c["closing"])}</p><p class="sign"><b>{e(c["signoff"])}</b></p></div>')
    css = (
        f'body{{margin:0;font-family:{font};font-size:11pt;color:#1a1a1a;line-height:1.45}}'
        f'.np-page{{padding:2.2cm 2.3cm}}h1{{font-size:{t["html_name_size"]};color:{t["html_heading"]};margin:0 0 4px}}'
        f'.muted{{color:{t["html_muted"]};font-size:{t["html_contact_size"]};margin:0 0 8px}}'
        f'hr{{border:0;border-top:1px solid {t["html_rule"]};margin:8px 0 18px}}p{{margin:0 0 10pt}}.sign{{margin-top:26pt}}'
    )
    return f"<!doctype html><html><head><meta charset=\"utf-8\"><style>{css}</style></head><body>{''.join(parts)}</body></html>"
