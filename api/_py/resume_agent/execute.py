"""Spec sections 11 and 16: run the generated script and check what it made.

`execute_python` is the tool contract from section 16. `validate_docx` checks
the file against the document acceptance tests in section 17. `render_check`
is the optional visual validation in 11.2; it only runs where LibreOffice is
installed and says so when it is not. `generate_document` is phases 3 and 4
together, including the correct-and-retry path for a failed run.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

from .filename import resume_filename
from .schema import clean_text, normalise, validate
from .script import build_script

SCRIPT_NAME = "build_resume.py"


def execute_python(code, expected_output_filename, workdir=None, timeout=60):
    """Write `code` to a working directory, run it with Python, report what happened."""
    workdir = workdir or tempfile.mkdtemp(prefix="resume-")
    os.makedirs(workdir, exist_ok=True)
    script_path = os.path.join(workdir, SCRIPT_NAME)
    with open(script_path, "w", encoding="utf-8") as f:
        f.write(code)
    # The child gets the same packages as this process (Vercel puts them on sys.path).
    env = {"PATH": os.environ.get("PATH", ""), "PYTHONPATH": os.pathsep.join(p for p in sys.path if p and os.path.isdir(p)),
           "PYTHONIOENCODING": "utf-8", "HOME": workdir, "LANG": "C.UTF-8"}
    try:
        proc = subprocess.run([sys.executable, script_path, workdir], cwd=workdir, env=env,
                              capture_output=True, text=True, timeout=timeout)
        exit_code, stdout, stderr = proc.returncode, proc.stdout, proc.stderr
    except subprocess.TimeoutExpired as e:
        exit_code, stdout, stderr = -1, e.stdout or "", f"Timed out after {timeout} seconds."
    files = []
    for name in sorted(os.listdir(workdir)):
        path = os.path.join(workdir, name)
        if name.endswith(".docx") and os.path.isfile(path):
            files.append({"filename": name, "path": path, "size_bytes": os.path.getsize(path)})
    return {"exit_code": exit_code, "stdout": (stdout or "")[-4000:], "stderr": (stderr or "")[-4000:],
            "files": files, "expected_output_filename": expected_output_filename, "workdir": workdir}


def _check(checks, name, ok, detail=""):
    checks.append({"name": name, "ok": bool(ok), "detail": detail})
    return ok


def _near(value, cm, tol=0.05):
    return value is not None and abs(value / 360000 - cm) <= tol


def validate_docx(path, data):
    """Open the document and check it against the spec. Returns {ok, checks}."""
    from docx import Document
    from docx.oxml.ns import qn

    checks = []
    if not _check(checks, "File exists", os.path.isfile(path), path):
        return {"ok": False, "checks": checks}
    if not _check(checks, "File is not empty", os.path.getsize(path) > 0, f"{os.path.getsize(path)} bytes"):
        return {"ok": False, "checks": checks}
    try:
        doc = Document(path)
        _check(checks, "Opens with python-docx (not corrupt)", True)
    except Exception as e:  # noqa: BLE001
        _check(checks, "Opens with python-docx (not corrupt)", False, repr(e))
        return {"ok": False, "checks": checks}

    body = doc.element.body
    paras = doc.paragraphs
    texts = [p.text for p in paras]
    alltext = "\n".join(texts)
    c = data["candidate"]

    s = doc.sections[0]
    _check(checks, "A4 page size", _near(s.page_width, 21.0) and _near(s.page_height, 29.7),
           f"{s.page_width / 360000:.2f} x {s.page_height / 360000:.2f} cm")
    _check(checks, "Portrait orientation", s.page_width < s.page_height)
    margins = [m / 360000 for m in (s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)]
    _check(checks, "Margins between 1.5 and 2 cm", all(1.49 <= m <= 2.01 for m in margins), ", ".join(f"{m:.2f}" for m in margins))
    _check(checks, "One section, no columns", len(doc.sections) == 1 and not body.findall(".//" + qn("w:cols") + "[@" + qn("w:num") + "]"))

    first = paras[0] if paras else None
    name_ok = first is not None and first.text == c["full_name"].upper() and first.alignment == 1
    if name_ok:
        r = first.runs[0]
        name_ok = bool(r.bold) and r.font.size is not None and r.font.size.pt == 22
    _check(checks, "Name: uppercase, centred, bold, 22 pt", name_ok, first.text if first is not None else "")
    contact_bits = [x for x in (c["location"], c["phone"], c["email"]) if x]
    if contact_bits:
        cp = next((p for p in paras[1:4] if all(b in p.text for b in contact_bits)), None)
        _check(checks, "Location, phone and email centred", cp is not None and cp.alignment == 1)

    _check(checks, "Candidate name present", c["full_name"].upper() in alltext)
    if data["professional_summary"]:
        _check(checks, "Professional summary present", data["professional_summary"] in alltext)
    headings = [p for p in paras if p._p.pPr is not None and p._p.pPr.find(qn("w:pBdr")) is not None]
    _check(checks, "At least one section heading", len(headings) > 0, ", ".join(h.text for h in headings))
    bordered = all(len(h._p.pPr.find(qn("w:pBdr")).findall(qn("w:bottom"))) == 1 for h in headings)
    _check(checks, "Section headings have one black bottom border", bordered and all(
        h._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")).get(qn("w:color")) == "000000" for h in headings))
    _check(checks, "No bordered boxes or shading", not body.findall(".//" + qn("w:shd")) and not body.findall(".//" + qn("w:pBdr") + "/" + qn("w:top")))
    for role in data["professional_experience"][:1]:
        _check(checks, "Employment content present", (role["job_title"] or role["employer"]) in alltext)

    _check(checks, "No tables", len(doc.tables) == 0)
    _check(checks, "No text boxes or drawings", not body.findall(".//" + qn("w:txbxContent")) and not body.findall(".//" + qn("w:drawing")))
    _check(checks, "No hyperlinks (no blue links)", not body.findall(".//" + qn("w:hyperlink")))
    hf_text = "".join(p.text for sec in doc.sections for part in (sec.header, sec.footer) for p in part.paragraphs)
    _check(checks, "Nothing in headers or footers", not hf_text.strip())
    _check(checks, "No manual page breaks", not body.findall(".//" + qn("w:br") + "[@" + qn("w:type") + "='page']"))
    _check(checks, "No empty paragraphs (no accidental blank space or pages)", all(t.strip() for t in texts))

    bullets = [p for p in paras if p.style.name == "List Bullet"]
    expect_bullets = bool(data["key_skills"] or any(r["achievements"] for r in data["professional_experience"]))
    if expect_bullets:
        _check(checks, "Bullets use the List Bullet style", bullets and not any(re.match(r"^\s*[•\-•●*]\s", p.text) for p in paras))

    bad_font, bad_colour, bad_size = [], [], []
    for p in paras:
        for r in p.runs:
            if not r.text.strip():
                continue
            fonts = r._element.rPr.find(qn("w:rFonts")) if r._element.rPr is not None else None
            if r.font.name != "Calibri" or fonts is None or fonts.get(qn("w:eastAsia")) != "Calibri":
                bad_font.append(r.text[:30])
            if r.font.color is None or r.font.color.rgb is None or str(r.font.color.rgb) != "000000":
                bad_colour.append(r.text[:30])
            size = r.font.size.pt if r.font.size else None
            if p is not first and not (size == 11 or (size == 12 and p.alignment == 1)):
                bad_size.append(f"{r.text[:30]} ({size})")
    _check(checks, "All text Calibri (Latin and East Asian)", not bad_font, "; ".join(bad_font[:3]))
    _check(checks, "All text black", not bad_colour, "; ".join(bad_colour[:3]))
    _check(checks, "Body text 11 pt", not bad_size, "; ".join(bad_size[:3]))
    normal = doc.styles["Normal"]
    _check(checks, "Normal and List Bullet styles set to Calibri 11 pt",
           normal.font.name == "Calibri" and normal.font.size.pt == 11 and doc.styles["List Bullet"].font.name == "Calibri")
    return {"ok": all(x["ok"] for x in checks), "checks": checks}


def render_check(path, timeout=60):
    """Optional visual validation (spec 11.2). Converts to PDF with LibreOffice.

    Returns {performed: False, reason} when LibreOffice is not available, so the
    caller never claims a visual check that did not happen.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {"performed": False, "reason": "Document rendering is not available here, so no visual check was done."}
    outdir = tempfile.mkdtemp(prefix="render-")
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", outdir, path],
                       capture_output=True, timeout=timeout, check=False,
                       env={**os.environ, "HOME": outdir})
        pdf = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
        if not os.path.isfile(pdf):
            return {"performed": False, "reason": "The renderer did not produce a PDF."}
        from pypdf import PdfReader
        reader = PdfReader(pdf)
        pages = [(pg.extract_text() or "").strip() for pg in reader.pages]
        problems = []
        if pages and not pages[0]:
            problems.append("The first page is blank.")
        problems += [f"Page {i + 1} is blank." for i, t in enumerate(pages) if i and not t]
        headings = {"PROFESSIONAL SUMMARY", "KEY SKILLS", "PROFESSIONAL EXPERIENCE", "EDUCATION AND TRAINING",
                    "AWARDS", "CERTIFICATIONS", "TECHNICAL SKILLS", "ADDITIONAL INFORMATION"}
        for i, t in enumerate(pages[:-1]):
            lines = [x.strip() for x in t.splitlines() if x.strip()]
            if lines and lines[-1] in headings:
                problems.append(f"The heading {lines[-1]} sits alone at the bottom of page {i + 1}.")
        return {"performed": True, "pages": len(pages), "problems": problems, "ok": not problems}
    except Exception as e:  # noqa: BLE001
        return {"performed": False, "reason": f"Rendering failed: {e!r}"}
    finally:
        shutil.rmtree(outdir, ignore_errors=True)


def _scrub(data):
    """Second-attempt correction: re-clean every string and drop anything odd."""
    def walk(v):
        if isinstance(v, dict):
            return {k: walk(x) for k, x in v.items()}
        if isinstance(v, list):
            return [walk(x) for x in v]
        return clean_text(v).encode("utf-8", "ignore").decode("utf-8") if isinstance(v, str) else v
    return normalise(walk(data))


def generate_document(data, fallback_title="", visual=True, keep_dir=False):
    """Phases 3 and 4: script, run, validate, and retry once with a corrected script.

    Returns a dict with `ok`, `filename`, `docx` (bytes, only when ok), `script`,
    `attempts` (one tool-contract result per run), `validation`, `visual` and
    `message` (the exact words to show the person).
    """
    data = normalise(data)
    problems = validate(data)
    if problems:
        return {"ok": False, "state": "EXECUTION_FAILED", "message": "The resume JSON is not valid: " + " ".join(problems),
                "attempts": [], "script": "", "filename": ""}
    filename = resume_filename(data, fallback_title)
    attempts, script, validation, result = [], "", None, None
    for attempt in (1, 2):
        if attempt == 2:
            data = _scrub(data)
        script = build_script(data, fallback_title)
        result = execute_python(script, filename)
        attempts.append({k: v for k, v in result.items() if k != "workdir"})
        path = os.path.join(result["workdir"], filename)
        if result["exit_code"] == 0 and os.path.isfile(path):
            validation = validate_docx(path, data)
            if validation["ok"]:
                break
        else:
            validation = {"ok": False, "checks": [{"name": "Script ran", "ok": False, "detail": (result["stderr"] or result["stdout"])[-600:]}]}
        if attempt == 1:
            shutil.rmtree(result["workdir"], ignore_errors=True)

    path = os.path.join(result["workdir"], filename)
    out = {"filename": filename, "script": script, "attempts": attempts, "validation": validation}
    if validation and validation["ok"]:
        with open(path, "rb") as f:
            out["docx"] = f.read()
        out["visual"] = render_check(path) if visual else {"performed": False, "reason": "Visual check not requested."}
        out.update(ok=True, state="COMPLETED", message=f"{filename} successfully generated.")
    else:
        failed = [x for x in (validation or {}).get("checks", []) if not x["ok"]]
        out.update(ok=False, state="EXECUTION_FAILED",
                   message="The Word document could not be generated and validated, so no file is being returned. "
                           + "; ".join(f"{x['name']}{': ' + x['detail'] if x['detail'] else ''}" for x in failed[:4]))
    if keep_dir:
        out["workdir"] = result["workdir"]
    else:
        shutil.rmtree(result["workdir"], ignore_errors=True)
    return out
